from __future__ import annotations

import datetime as dt

import pytest
from openpyxl import load_workbook

from clasificador.analysis.analyzer import recalcular_categorias
from clasificador.analysis.rules_config import DEFAULT_COMPILED_RULES
from clasificador.analysis.schema import (
    CategoriaDef,
    CondicionIndicador,
    EsquemaClasificacion,
    IndicadorDef,
    ReglaDecision,
)
from clasificador.analysis.schema_engine import clasificar_por_esquema, compilar_esquema
from clasificador.conversation_view import (
    ConversationTranscript,
    TranscriptMessage,
    category_counts,
    export_conversations,
    subcategorias_disponibles,
    subcategory_counts,
)
from clasificador.db.engine import create_all_tables, create_db_engine, make_session_factory
from clasificador.db.models import Campana, Conversation, ConversationAnalysis
from clasificador.db.repository import start_import
from clasificador.reports import build_conversations_xlsx

T0 = dt.datetime(2026, 9, 1, 9, 0, 0)


def _esquema() -> EsquemaClasificacion:
    """Como el de Acme: 2 categorías + SIN_RESPUESTA oculta, y el motivo del
    interés como subcategoría (solo aplica a INTERESADO)."""
    return EsquemaClasificacion(
        categorias=(
            CategoriaDef("SIN_RESPUESTA", "Sin respuesta", ocultar_en_dashboard=True),
            CategoriaDef("INTERESADO", "Interesado"),
            CategoriaDef("NO_INTERESADO", "No interesado"),
        ),
        indicadores=(
            IndicadorDef(
                "rechazo", "Rechazo", "regla", condiciones=(CondicionIndicador("boton", "no estoy interesado"),)
            ),
            IndicadorDef("interes", "Interés", "ia", descripcion_ia="El cliente muestra interés"),
            IndicadorDef("vehiculos", "Vehículos", "ia", descripcion_ia="Interés en motos o carros"),
        ),
        arbol_decision=(
            ReglaDecision({"rechazo": True}, "NO_INTERESADO"),
            ReglaDecision({"vehiculos": True}, "INTERESADO"),
            ReglaDecision({"interes": True}, "INTERESADO"),
            ReglaDecision({}, "NO_INTERESADO"),
        ),
        categoria_sin_respuesta="SIN_RESPUESTA",
        subcategorias=(CategoriaDef("VEH", "Vehículos"), CategoriaDef("SIN_DETALLE", "Sin detalle")),
        arbol_subcategoria=(ReglaDecision({"vehiculos": True}, "VEH"), ReglaDecision({}, "SIN_DETALLE")),
        subcategoria_aplica_a=("INTERESADO",),
        subcategoria_etiqueta="Motivo del interés",
    )


def _transcript(*pares) -> ConversationTranscript:
    msgs = [
        TranscriptMessage(sequence_number=i, timestamp=T0, sender=s, sender_label=s, message_type="TEXT", message=t)
        for i, (s, t) in enumerate(pares, start=1)
    ]
    return ConversationTranscript(
        session_id="sid",
        customer_phone="57",
        customer_name="P",
        started_at=T0,
        ended_at=T0,
        conn_id_count=1,
        messages=msgs,
    )


def test_roundtrip_y_validar_con_subcategorias():
    esquema = _esquema()
    assert esquema.validar() is None
    restaurado = EsquemaClasificacion.from_json(esquema.to_json())
    assert restaurado == esquema
    assert restaurado.categorias[0].ocultar_en_dashboard is True


def test_validar_detecta_subcategoria_no_declarada_y_aplica_a_invalido():
    base = _esquema()
    malo = EsquemaClasificacion.from_json(
        {**base.to_json(), "arbol_subcategoria": [{"condicion": {}, "categoria": "NO_EXISTE"}]}
    )
    assert malo.validar() is not None
    malo2 = EsquemaClasificacion.from_json({**base.to_json(), "subcategoria_aplica_a": ["NO_EXISTE"]})
    assert malo2.validar() is not None


def test_esquema_sin_subcategorias_sigue_igual():
    sin = EsquemaClasificacion.from_json(
        {
            "categorias": [{"id": "A", "label": "A"}],
            "indicadores": [{"id": "x", "label": "X", "tipo": "ia", "descripcion_ia": "d"}],
            "arbol_decision": [{"condicion": {}, "categoria": "A"}],
            "categoria_sin_respuesta": "A",
        }
    )
    assert sin.validar() is None
    assert sin.subcategorias == ()


def test_motor_rechazo_explicito_no_tiene_subcategoria_ni_llama_a_ia():
    compilado = compilar_esquema(_esquema())
    t = _transcript(("BOT", "hola"), ("CUSTOMER", "No estoy interesado"))
    out = clasificar_por_esquema(t, compilado, DEFAULT_COMPILED_RULES)
    assert (out.categoria, out.subcategoria, out.requiere_ia) == ("NO_INTERESADO", None, False)


def test_motor_pide_ia_cuando_la_subcategoria_depende_de_una_senal_ia():
    """La categoría INTERESADO ya salía igual con cualquier señal de IA, pero la
    subcategoría (motivo) depende de `vehiculos` — hay que preguntarle al modelo."""
    compilado = compilar_esquema(_esquema())
    t = _transcript(("BOT", "hola"), ("CUSTOMER", "quiero una moto financiada por favor"))
    out = clasificar_por_esquema(t, compilado, DEFAULT_COMPILED_RULES)
    assert out.requiere_ia is True


def test_motor_sin_respuesta():
    compilado = compilar_esquema(_esquema())
    out = clasificar_por_esquema(_transcript(("BOT", "hola")), compilado, DEFAULT_COMPILED_RULES)
    assert (out.categoria, out.subcategoria, out.requiere_ia) == ("SIN_RESPUESTA", None, False)


@pytest.fixture
def session_factory():
    engine = create_db_engine("sqlite:///:memory:")
    create_all_tables(engine)
    return make_session_factory(engine)


def _conv(session, imp_id, sid, *, senales, categoria, subcategoria=None):
    conv = Conversation(
        session_id=sid,
        session_pattern="REAL",
        campana_id=1,
        customer_phone="57",
        customer_name="C",
        started_at=T0,
        ended_at=T0,
        message_count=1,
        conn_id_count=1,
        iniciada_por="CUSTOMER",
        contiene_array=False,
        first_import_id=imp_id,
        last_import_id=imp_id,
    )
    session.add(conv)
    session.flush()
    session.add(
        ConversationAnalysis(
            conversation_id=conv.id,
            categoria_final=categoria,
            subcategoria=subcategoria,
            senales=senales,
            metodo="RULE",
            analyzed_at=T0,
        )
    )


def _armar(session_factory):
    with session_factory() as session:
        session.add(
            Campana(id_campana=1, nombre_campana="Acme", configuracion={"esquema_clasificacion": _esquema().to_json()})
        )
        session.flush()
        imp = start_import(session, "f.xlsx", "h", rows_received=3)
        # categorías "viejas" (7 categorías) que el recálculo debe llevar a 2 niveles
        _conv(session, imp.id, "a", senales={"rechazo": False, "interes": True, "vehiculos": True}, categoria="OTRA")
        _conv(session, imp.id, "b", senales={"rechazo": False, "interes": True, "vehiculos": False}, categoria="OTRA")
        _conv(session, imp.id, "c", senales={"rechazo": True, "interes": True, "vehiculos": True}, categoria="OTRA")
        _conv(session, imp.id, "d", senales={}, categoria="OTRA")
        session.commit()


def test_recalcular_categorias_sin_ia_y_conteos(session_factory):
    _armar(session_factory)
    rep = recalcular_categorias(session_factory)
    assert rep.revisadas == 4 and rep.cambiadas == 4

    with session_factory() as session:
        assert category_counts(session) == {"INTERESADO": 2, "NO_INTERESADO": 1, "SIN_RESPUESTA": 1}
        assert subcategory_counts(session) == {"VEH": 1, "SIN_DETALLE": 1}
        cat = subcategorias_disponibles(session)
        assert cat["etiqueta"] == "Motivo del interés"
        assert [i["id"] for i in cat["items"]] == ["VEH", "SIN_DETALLE"]

    # idempotente: una segunda pasada no cambia nada
    assert recalcular_categorias(session_factory).cambiadas == 0


def test_excel_incluye_columna_de_subcategoria(session_factory):
    _armar(session_factory)
    recalcular_categorias(session_factory)
    with session_factory() as session:
        buffer = build_conversations_xlsx(export_conversations(session))
    ws = load_workbook(buffer).active
    encabezados = [c.value for c in ws[1]]
    assert "Motivo del interés" in encabezados
    idx = encabezados.index("Motivo del interés")
    valores = {r[0].value: r[idx].value for r in ws.iter_rows(min_row=2)}
    assert valores["a"] == "Vehículos"
    assert valores["b"] == "Sin detalle"
    assert valores["c"] in (None, "")  # No interesado: no aplica
