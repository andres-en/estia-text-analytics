from __future__ import annotations

import datetime as dt

import pytest
from openpyxl import load_workbook

from clasificador.conversation_view import export_conversations
from clasificador.db.engine import create_all_tables, create_db_engine, make_session_factory
from clasificador.db.models import Campana, Conversation, ConversationAnalysis
from clasificador.db.repository import start_import
from clasificador.reports import build_conversations_xlsx

T0 = dt.datetime(2026, 9, 1, 9, 0, 0)


@pytest.fixture
def session_factory():
    engine = create_db_engine("sqlite:///:memory:")
    create_all_tables(engine)
    return make_session_factory(engine)


def _crear_conversacion(session, imp_id, campana_id, session_id, *, senales=None, categoria="X", metodo="RULE"):
    conv = Conversation(
        session_id=session_id,
        session_pattern="REAL",
        campana_id=campana_id,
        customer_phone="573000000000",
        customer_name="Cliente",
        started_at=T0,
        ended_at=T0,
        message_count=1,
        conn_id_count=1,
        iniciada_por="CUSTOMER",
        contiene_array=False,
        first_import_id=imp_id,
        last_import_id=imp_id,
    )
    session.add(conv)
    session.flush()
    if senales is not None:
        session.add(
            ConversationAnalysis(
                conversation_id=conv.id,
                categoria_final=categoria,
                senales=senales,
                metodo=metodo,
                analyzed_at=T0,
            )
        )
    return conv


def test_excel_arma_columnas_legadas_y_del_esquema_propio(session_factory):
    """Bug real: el reporte tenía 6 columnas fijas leyendo columnas legadas de
    la base — una campaña con EsquemaClasificacion propio (ver
    analysis/schema.py) quedaba con esas columnas siempre en blanco. Ahora las
    columnas se arman a partir de `senales`, así que la campaña con esquema
    propio muestra SUS indicadores, no los de Acme."""
    with session_factory() as session:
        session.add(Campana(id_campana=1, nombre_campana="Acme (legado)", configuracion=None))
        session.add(
            Campana(
                id_campana=2,
                nombre_campana="Cobranza",
                configuracion={
                    "esquema_clasificacion": {
                        "categorias": [{"id": "X", "label": "X"}],
                        "indicadores": [
                            {
                                "id": "interes_vehiculos",
                                "label": "Interés en vehículos",
                                "tipo": "ia",
                                "descripcion_ia": "algo",
                            }
                        ],
                        "arbol_decision": [{"condicion": {}, "categoria": "X"}],
                        "categoria_sin_respuesta": "X",
                    }
                },
            )
        )
        session.flush()

        imp = start_import(session, "f.xlsx", "h1", rows_received=2)

        # Conversación de una campaña sin esquema propio (motor fijo, las 6
        # señales legadas de siempre).
        _crear_conversacion(
            session,
            imp.id,
            1,
            "sid-legado",
            senales={
                "interes_declarado": True,
                "interes_confirmado": False,
                "busco_informacion": True,
                "rechazo_explicito": False,
                "bot_no_resolvio": False,
                "es_autorespondedor": False,
            },
        )
        # Conversación de una campaña con esquema propio.
        _crear_conversacion(session, imp.id, 2, "sid-esquema", senales={"interes_vehiculos": True})

        session.commit()

        export = export_conversations(session)

    assert export.indicador_labels["interes_declarado"] == "Declared interest"
    assert export.indicador_labels["interes_vehiculos"] == "Interés en vehículos"

    buffer = build_conversations_xlsx(export)
    wb = load_workbook(buffer)
    ws = wb.active
    encabezados = [c.value for c in ws[1]]

    # Los 6 legados van primero (en su orden de siempre), la columna del
    # esquema propio se agrega después.
    idx_legado = encabezados.index("Declared interest")
    idx_esquema = encabezados.index("Interés en vehículos")
    assert idx_esquema > idx_legado
    # DEFAULT_ESQUEMA (analysis/schema.py) es la fuente de las etiquetas para
    # una campaña sin esquema propio — mismo criterio que ya usa la resolución
    # de categorías (_resolver_categoria_label).
    assert encabezados[idx_legado : idx_legado + 6] == [
        "Declared interest",
        "Confirmed interest",
        "Sought information",
        "Explicit rejection",
        "Bot could not help",
        "Business auto-reply",
    ]

    filas = {row[0].value: row for row in ws.iter_rows(min_row=2)}
    fila_legado = filas["sid-legado"]
    assert fila_legado[idx_legado].value == "Yes"  # interes_declarado
    assert fila_legado[idx_legado + 2].value == "Yes"  # busco_informacion
    assert fila_legado[idx_esquema].value is None  # esa campaña no tiene este indicador

    fila_esquema = filas["sid-esquema"]
    assert fila_esquema[idx_esquema].value == "Yes"
    assert fila_esquema[idx_legado].value is None  # esa campaña no tiene los indicadores legados


def test_excel_conversacion_sin_analizar_deja_columnas_en_blanco(session_factory):
    with session_factory() as session:
        session.add(Campana(id_campana=1, nombre_campana="Acme", configuracion=None))
        session.flush()
        imp = start_import(session, "f.xlsx", "h1", rows_received=1)
        _crear_conversacion(session, imp.id, 1, "sid-legado", senales={"interes_declarado": True})
        _crear_conversacion(session, imp.id, 1, "sid-sin-analizar", senales=None)
        session.commit()

        export = export_conversations(session)

    buffer = build_conversations_xlsx(export)
    wb = load_workbook(buffer)
    ws = wb.active
    encabezados = [c.value for c in ws[1]]
    idx_cat = encabezados.index("Category")
    idx_ind = encabezados.index("Declared interest")

    filas = {row[0].value: row for row in ws.iter_rows(min_row=2)}
    assert filas["sid-sin-analizar"][idx_cat].value == "Not analyzed"
    assert filas["sid-sin-analizar"][idx_ind].value is None
