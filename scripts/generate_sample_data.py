"""Generate a fully synthetic CDR Excel file for local development and tests.

The real CDR exports contain customer data, so they are never committed. This
script builds a file with the exact same layout (metadata rows + header on row 8
+ the 12 CDR columns) filled with fictitious customers, phones and messages for
a fictitious credit-line product ("Acme").

It deliberately includes every data-quality edge case the ingestion pipeline
handles, so the integration tests exercise all of them:

- REAL conversations and proactive (broadcast) messages with both broadcast
  SESSION_ID formats (33-digit numeric and dotted).
- A multi-segment conversation: 4 CONN_IDs and 3 `end_exit_chat` closures.
- Adjacent duplicate customer rows emitted by the CDR (must be skipped).
- Broadcast SESSION_ID collisions between different phones (must warn).
- Conversations covering every classification path: confirmed interest,
  explicit rejection + survey, courtesy-only replies, company auto-responders,
  "bot could not help" answers, and free-text "gray zone" questions.

Usage:
    python scripts/generate_sample_data.py [--output data/samples/sample_cdr.xlsx]

The output is deterministic (fixed seed): re-running it produces the same data,
which the test baselines rely on.
"""

from __future__ import annotations

import argparse
import datetime as dt
import random
from pathlib import Path

import openpyxl

DEFAULT_OUTPUT = Path(__file__).resolve().parent.parent / "data" / "samples" / "sample_cdr.xlsx"

HEADERS = [
    "SESSION_ID",
    "CHANNEL",
    "RP_ID",
    "DATE",
    "CUSTOMER_NAME",
    "CUSTOMER_EMAIL",
    "CUSTOMER_PHONE",
    "CUSTOMER_QUERY",
    "ROUTING_ANSWER",
    "CUSTOMER_ID",
    "CONN_ID",
    "BSUID",
]

TENANT = "botco-4402-acme_venta"
SESSION_SUFFIX = "1001"
RP_ID = 50001
DAY_1 = dt.datetime(2026, 8, 1)
DAY_2 = dt.datetime(2026, 8, 2)

# Well-known conversations referenced by the tests.
FIRST_PHONE = 573001000001  # first data row of the file
MULTI_SEGMENT_PHONE = 573001000002  # 4 CONN_IDs, 3 closures, 30 messages
SEARCH_PHONE = 573001000003  # used by the "search by phone" API test

FIRST_NAMES = [
    "Ana",
    "Luis",
    "Carlos",
    "Maria",
    "Jorge",
    "Paula",
    "Andres",
    "Laura",
    "Diana",
    "Camilo",
    "Sofia",
    "Julian",
    "Valentina",
    "Mateo",
    "Natalia",
    "Sergio",
    "Carolina",
    "Daniel",
    "Marta",
    "Oscar",
]

# --- Bot templates -------------------------------------------------------------
OFFER = (
    "¡Qué buena noticia! Con tu cupo Acme puedes financiar una gran variedad de productos, "
    "como electrodomésticos, tecnología, muebles, motos y materiales de construcción. "
    "¿Tienes en mente algún producto que quieras adquirir?"
)
GREETING = (
    "¡Hola! Es un gusto saludarte. Soy tu asesor digital de Acme y estoy aquí para ayudarte "
    "a aprovechar al máximo tu cupo. ¿Tienes en mente algún producto que quieras adquirir?"
)
ASK_PARTNER = (
    "¡Excelente elección! ¿Con cuál de nuestros aliados te gustaría comprarlo? Tenemos tiendas en toda la ciudad."
)
ASK_ADVISOR = "¿Te gustaría que un asesor te contacte para completar la compra? Responde Sí o No."
ASK_CITY = "Perfecto. ¿Me confirmas en qué ciudad te encuentras?"
REGISTERED = (
    "¡Listo! Hemos registrado tu interés y un asesor te contactará en las próximas horas "
    "para ayudarte a completar tu compra."
)
REJECTION_SURVEY = (
    "Lamentamos que por ahora no desees utilizar tu cupo Acme. Para ayudarnos a mejorar, "
    "cuéntanos el principal motivo: 1️⃣ No necesito comprar en este momento. 2️⃣ Prefiero no "
    "endeudarme. 3️⃣ No conozco los aliados. 4️⃣ Otro motivo."
)
SURVEY_THANKS = "¡Gracias por contarnos! Recuerda que tu cupo sigue disponible cuando lo necesites."
INTEREST_INFO = (
    "La tasa de interés depende del plazo que elijas y se informa antes de confirmar la compra. "
    "Puedes diferir tu compra hasta en 60 meses."
)
PARTNERS_INFO = "Puedes usar tu cupo en cualquiera de nuestros aliados: tiendas de tecnología, muebles y motos."
CASH_INFO = "Tu cupo Acme no entrega dinero en efectivo: se usa para comprar productos con nuestros aliados."
BOT_CANT_HELP = "Lo siento, no tengo acceso a esa información en este momento. Un asesor podrá ayudarte."
YOU_ARE_WELCOME = "¡Con mucho gusto! Aquí estaré si necesitas algo más."
DELAY_APOLOGY = "Lamento la demora. Ya escalé tu caso para que te atiendan lo antes posible."
REMINDERS = [
    "¡Hola! Espero que estés teniendo un buen día. Te escribo para que no olvides que dejamos pendiente tu solicitud.",
    "¡Hola! Paso por aquí para que este chat no se nos quede en el olvido. 📝",
    "¡Hola, hola! Un mensajito rápido para que no te olvides de responderme por aquí cuando puedas. 😉",
]
END = "end_exit_chat"
PROACTIVE = (
    "👋 Hola, *{name}*. 💰 Hoy tienes disponible un cupo Acme de *${cupo}*, que puedes utilizar con "
    "nuestros aliados y pagar cómodamente a través de tu factura de servicios públicos. 🛍️ "
    "¿Quieres usarlo? Elige una opción: Quiero usar mi cupo / Hablar con un asesor / No estoy interesado"
)

AUTORESPONDERS = [
    "Gracias por comunicarte con Ferretería El Tornillo. En este momento no podemos responder, pero lo haremos lo antes posible.",
    "¡Bienvenido(a) a Panadería La Espiga! Nuestro horario de atención es de lunes a sábado de 7am a 7pm.",
    "Gracias por comunicarse con Taller Rueda Libre. ¿Cómo podemos ayudarte?",
]
PRODUCTS = ["Moto", "Nevera", "Un televisor", "Celular", "Muebles", "Lavadora", "Computador"]
PARTNERS = ["Tienda Central", "Motos del Norte", "ElectroHogar", "Muebles Nova"]
CITIES = ["Barranquilla", "Cartagena", "Santa Marta", "Montería"]
QUESTIONS = [
    ("Que intereses manejan", INTEREST_INFO),
    ("Donde puedo usarlo", PARTNERS_INFO),
    ("Efectivo", CASH_INFO),
    ("A cuantos meses puedo diferir", INTEREST_INFO),
]
UNANSWERABLE = [
    "Cuanto debo en la factura de este mes",
    "Me pueden decir el saldo exacto de mi cupo",
]


class Clock:
    """Hands out strictly increasing timestamps within one conversation."""

    def __init__(self, start: dt.datetime, rng: random.Random):
        self.now = start
        self.rng = rng

    def tick(self, min_s: int = 20, max_s: int = 240) -> dt.datetime:
        self.now += dt.timedelta(seconds=self.rng.randint(min_s, max_s))
        return self.now


def real_session_id(phone: int) -> str:
    return f"{phone}.{TENANT}.{SESSION_SUFFIX}"


def real_row(sid, phone, date, conn_id, query=None, answer=None) -> list:
    return [sid, "whatsapp", RP_ID, date, "-", "-", phone, query or "-", answer or "-", "-", conn_id, "---"]


def build_conversation(kind: str, phone: int, start: dt.datetime, rng: random.Random, conn_ids) -> list[list]:
    """Return the CDR rows for one REAL conversation of the given kind."""
    sid = real_session_id(phone)
    clock = Clock(start, rng)
    conn = next(conn_ids)
    rows: list[list] = []

    def add(query=None, answer=None, same_time=False, conn_id=None):
        date = clock.now if same_time else clock.tick()
        rows.append(real_row(sid, phone, date, conn_id or conn, query, answer))

    if kind == "interested":
        add("Quiero usar mi cupo", OFFER)
        add(rng.choice(PRODUCTS), ASK_PARTNER)
        add(rng.choice(PARTNERS), ASK_ADVISOR)
        add("Sí", ASK_CITY)
        add(rng.choice(CITIES), REGISTERED)
        add(answer=END)
    elif kind == "rejected":
        add("No estoy interesado", REJECTION_SURVEY)
        add(str(rng.randint(1, 4)), SURVEY_THANKS)
        add(answer=END)
    elif kind == "rejected_with_duplicate":
        # The CDR sometimes emits the same customer message twice (same DATE +
        # CONN_ID + text): the second copy must not become a new message.
        add("No estoy interesado", REJECTION_SURVEY)
        add("No estoy interesado", same_time=True)
        add(str(rng.randint(1, 4)), SURVEY_THANKS)
        add(answer=END)
    elif kind == "courtesy_only":
        add("Hola", GREETING)
        add("Gracias", YOU_ARE_WELCOME)
        add(answer=rng.choice(REMINDERS))
    elif kind == "autoresponder":
        add(rng.choice(AUTORESPONDERS), GREETING)
        add(answer=rng.choice(REMINDERS))
        add(answer=END)
    elif kind == "gray_zone":
        question, info = rng.choice(QUESTIONS)
        add("Quiero usar mi cupo", OFFER)
        add(question, info)
        add("Ok gracias", YOU_ARE_WELCOME)
        add(answer=END)
    elif kind == "bot_could_not_help":
        add("Hola", GREETING)
        add(rng.choice(UNANSWERABLE), BOT_CANT_HELP)
        add(answer=rng.choice(REMINDERS))
        add(answer=END)
    elif kind == "multi_segment":
        # 4 segments (CONN_IDs), 3 closures and 30 messages in total; "Sí" is sent
        # once per closed segment, and duplicated by the CDR in two of them.
        add("Quiero usar mi cupo", OFFER)
        add("Moto", ASK_ADVISOR)
        add("Sí", ASK_CITY)
        add("Sí", same_time=True)
        add(answer=REMINDERS[0])
        add(answer=END)

        conn = next(conn_ids)
        add("Hola", GREETING, conn_id=conn)
        add("Quiero hablar con alguien de la moto", ASK_ADVISOR, conn_id=conn)
        add("Sí", ASK_CITY, conn_id=conn)
        add(answer=REMINDERS[1], conn_id=conn)
        add(answer=END, conn_id=conn)

        conn = next(conn_ids)
        add("Buenas tardes", GREETING, conn_id=conn)
        add("Ya les había escrito, sigo esperando", DELAY_APOLOGY, conn_id=conn)
        add(answer=REMINDERS[2], conn_id=conn)

        conn = next(conn_ids)
        add("Hola", GREETING, conn_id=conn)
        add("Sí", ASK_CITY, conn_id=conn)
        add("Sí", same_time=True, conn_id=conn)
        add("Cartagena", REGISTERED, conn_id=conn)
        add("Gracias", YOU_ARE_WELCOME, conn_id=conn)
        add(answer=END, conn_id=conn)
    else:
        raise ValueError(f"Unknown conversation kind: {kind}")
    return rows


def broadcast_session_id(rng: random.Random) -> str:
    if rng.random() < 0.1:
        return ".".join(str(rng.randint(10 ** (n - 1), 10**n - 1)) for n in (4, 12, 14, 2))
    return str(rng.randint(10**32, 10**33 - 1))


def broadcast_row(sid, phone, name, date, rng) -> list:
    cupo = rng.randrange(1_500_000, 6_000_000, 5_000)
    return [
        sid,
        "whatsapp",
        RP_ID,
        date,
        name,
        "-",
        phone,
        PROACTIVE.format(name=name, cupo=cupo),
        "-",
        str(rng.randint(10_000_000, 99_999_999)),
        "-",
        "---",
    ]


def generate(seed: int = 42) -> list[list]:
    rng = random.Random(seed)
    conn_ids = iter(range(900_000_001, 999_999_999))

    plan = (
        ["interested"] * 12
        + ["rejected"] * 16
        + ["rejected_with_duplicate"] * 4
        + ["courtesy_only"] * 8
        + ["autoresponder"] * 5
        + ["gray_zone"] * 10
        + ["bot_could_not_help"] * 4
    )
    rng.shuffle(plan)
    # Pinned conversations first, so tests can rely on them.
    plan = ["interested", "multi_segment", "interested"] + plan

    rows: list[list] = []
    phones: list[int] = []
    for i, kind in enumerate(plan):
        phone = FIRST_PHONE + i  # 573001000001, ...002 (multi-segment), ...003 (search)
        phones.append(phone)
        day = DAY_1 if i % 2 == 0 else DAY_2
        start = day.replace(hour=rng.randint(10, 20), minute=rng.randint(0, 59), second=rng.randint(0, 59))
        rows.extend(build_conversation(kind, phone, start, rng, conn_ids))

    # Proactive messages: every customer above plus many who never replied.
    broadcast_phones = phones + [573100000000 + n for n in range(1, 896 - len(phones) + 1)]
    broadcast_sids = []
    for phone in broadcast_phones:
        day = DAY_1 if rng.random() < 0.5 else DAY_2
        date = day.replace(hour=9, minute=rng.randint(0, 59), second=rng.randint(0, 59))
        sid = broadcast_session_id(rng)
        broadcast_sids.append(sid)
        rows.append(broadcast_row(sid, phone, rng.choice(FIRST_NAMES), date, rng))

    # 4 SESSION_ID collisions between different phones (seen in real exports).
    for k in range(4):
        date = DAY_2.replace(hour=11, minute=k, second=0)
        rows.append(broadcast_row(broadcast_sids[-1 - k], 573200000000 + k, rng.choice(FIRST_NAMES), date, rng))

    return rows


def write_workbook(rows: list[list], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HistIVR20260803-090805"
    ws["A3"] = "9. CDR Routing interacciones y chatbots"
    ws["A5"] = "REPORT: 9. CDR Routing interacciones y chatbots (SYNTHETIC SAMPLE)"
    ws["A6"] = "INI DATE: 2026/08/01"
    ws["A7"] = "END DATE: 2026/08/02"
    for col, header in enumerate(HEADERS, start=1):
        ws.cell(row=8, column=col, value=header)
    for r, row in enumerate(rows, start=9):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    wb.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args()

    rows = generate(args.seed)
    write_workbook(rows, args.output)
    print(f"Wrote {len(rows)} rows to {args.output}")


if __name__ == "__main__":
    main()
